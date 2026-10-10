"""Document service for RAG functionality."""

import hashlib
from pathlib import Path
from typing import Any
from uuid import uuid4

import aiofiles
from sqlalchemy import select, text
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.config import settings
from app.models.document import Document, DocumentChunk, DocumentStatus, DocumentType


class DocumentService:
    """Service for document processing and RAG."""

    def __init__(self, db: AsyncSession):
        self.db = db
        from langchain_text_splitters import RecursiveCharacterTextSplitter
        from langchain_openai import OpenAIEmbeddings

        self.embeddings = OpenAIEmbeddings(
            model=settings.EMBEDDING_MODEL,
            api_key=settings.OPENAI_API_KEY,
            dimensions=settings.EMBEDDING_DIMENSIONS,
        )
        self.text_splitter = RecursiveCharacterTextSplitter(
            chunk_size=settings.CHUNK_SIZE,
            chunk_overlap=settings.CHUNK_OVERLAP,
            length_function=len,
            separators=["\n\n", "\n", ". ", " ", ""],
        )

    def _get_document_type(self, filename: str) -> DocumentType:
        """Determine document type from filename."""
        ext = Path(filename).suffix.lower()
        type_map = {
            ".pdf": DocumentType.PDF,
            ".docx": DocumentType.DOCX,
            ".doc": DocumentType.DOCX,
            ".txt": DocumentType.TXT,
            ".md": DocumentType.MD,
            ".csv": DocumentType.CSV,
            ".xlsx": DocumentType.XLSX,
            ".xls": DocumentType.XLSX,
            ".json": DocumentType.JSON,
            ".html": DocumentType.HTML,
            ".htm": DocumentType.HTML,
            ".png": DocumentType.IMAGE,
            ".jpg": DocumentType.IMAGE,
            ".jpeg": DocumentType.IMAGE,
        }
        return type_map.get(ext, DocumentType.TXT)

    def _compute_content_hash(self, content: bytes) -> str:
        """Compute SHA-256 hash of content."""
        return hashlib.sha256(content).hexdigest()

    async def upload_document(
        self,
        user_id: str,
        filename: str,
        content: bytes,
        mime_type: str,
    ) -> Document:
        """Upload a document."""
        # Generate unique filename
        ext = Path(filename).suffix
        unique_filename = f"{uuid4()}{ext}"
        file_path = Path(settings.UPLOAD_DIR) / user_id / unique_filename

        # Ensure directory exists
        file_path.parent.mkdir(parents=True, exist_ok=True)

        # Save file
        async with aiofiles.open(file_path, "wb") as f:
            await f.write(content)

        # Create document record
        document = Document(
            user_id=user_id,
            filename=unique_filename,
            original_filename=filename,
            file_path=str(file_path),
            file_size=len(content),
            mime_type=mime_type,
            document_type=self._get_document_type(filename).value,
            content_hash=self._compute_content_hash(content),
            status=DocumentStatus.PENDING.value,
        )
        self.db.add(document)
        await self.db.flush()

        return document

    async def reprocess_document(self, document: Document) -> Document:
        """Retry processing a failed document."""
        if document.status != DocumentStatus.FAILED.value:
            raise ValueError(
                f"Only failed documents can be reprocessed, current status: {document.status}"
            )

        # Clear old chunks before reprocessing
        from sqlalchemy import delete
        await self.db.execute(
            delete(DocumentChunk).where(DocumentChunk.document_id == document.id)
        )
        document.error_message = None
        document.chunk_count = 0
        document.status = DocumentStatus.PENDING.value
        await self.db.flush()

        return await self.process_document(document)

    async def process_document(self, document: Document) -> Document:
        """Process a document: extract text, chunk, and embed."""
        if document.status not in (
            DocumentStatus.PENDING.value,
            DocumentStatus.FAILED.value,
        ):
            raise ValueError(
                f"Document must be in pending or failed state to process, "
                f"current status: {document.status}"
            )
        try:
            document.status = DocumentStatus.PROCESSING.value
            await self.db.flush()

            # Read file content
            async with aiofiles.open(document.file_path, "rb") as f:
                content = await f.read()

            # Extract text based on document type
            text_content = await self._extract_text(document, content)

            # Split into chunks
            chunks = self.text_splitter.split_text(text_content)

            # Generate embeddings
            embeddings = await self.embeddings.aembed_documents(chunks)

            # Create chunk records
            for i, (chunk_text, embedding) in enumerate(zip(chunks, embeddings)):
                chunk = DocumentChunk(
                    document_id=document.id,
                    content=chunk_text,
                    chunk_index=i,
                    embedding=embedding,
                    token_count=len(chunk_text.split()),  # Approximate
                )
                self.db.add(chunk)

            document.chunk_count = len(chunks)
            document.status = DocumentStatus.COMPLETED.value
            await self.db.flush()

        except Exception as e:
            document.status = DocumentStatus.FAILED.value
            document.error_message = str(e)
            await self.db.flush()
            raise

        return document

    async def _extract_text(self, document: Document, content: bytes) -> str:
        """Extract text from document based on type."""
        doc_type = DocumentType(document.document_type)

        if doc_type == DocumentType.PDF:
            from pypdf import PdfReader
            from io import BytesIO

            reader = PdfReader(BytesIO(content))
            return "\n\n".join(page.extract_text() for page in reader.pages)

        elif doc_type == DocumentType.DOCX:
            from docx import Document as DocxDocument
            from io import BytesIO

            doc = DocxDocument(BytesIO(content))
            return "\n\n".join(para.text for para in doc.paragraphs)

        elif doc_type in (DocumentType.TXT, DocumentType.MD):
            return content.decode("utf-8")

        elif doc_type == DocumentType.HTML:
            from bs4 import BeautifulSoup

            soup = BeautifulSoup(content, "html.parser")
            return soup.get_text(separator="\n")

        elif doc_type == DocumentType.JSON:
            import json

            data = json.loads(content)
            return json.dumps(data, indent=2)

        else:
            return content.decode("utf-8", errors="ignore")

    async def search_documents(
        self,
        user_id: str,
        query: str,
        limit: int = 10,
        similarity_threshold: float = 0.7,
    ) -> list[dict[str, Any]]:
        """Search documents using vector similarity."""
        # Generate query embedding
        query_embedding = await self.embeddings.aembed_query(query)

        # Vector similarity search
        result = await self.db.execute(
            text("""
                SELECT
                    dc.id,
                    dc.document_id,
                    dc.content,
                    dc.chunk_index,
                    d.original_filename,
                    1 - (dc.embedding <=> :embedding::vector) as similarity
                FROM document_chunks dc
                JOIN documents d ON dc.document_id = d.id
                WHERE d.user_id = :user_id
                    AND d.status = 'completed'
                    AND 1 - (dc.embedding <=> :embedding::vector) > :threshold
                ORDER BY dc.embedding <=> :embedding::vector
                LIMIT :limit
            """),
            {
                "embedding": query_embedding,
                "user_id": user_id,
                "threshold": similarity_threshold,
                "limit": limit,
            },
        )

        rows = result.fetchall()
        return [
            {
                "id": str(row.id),
                "document_id": str(row.document_id),
                "content": row.content,
                "chunk_index": row.chunk_index,
                "filename": row.original_filename,
                "similarity": float(row.similarity),
            }
            for row in rows
        ]

    async def get_document(self, document_id: str, user_id: str) -> Document | None:
        """Get a document by ID."""
        result = await self.db.execute(
            select(Document).where(
                Document.id == document_id,
                Document.user_id == user_id,
            )
        )
        return result.scalar_one_or_none()

    async def list_documents(
        self,
        user_id: str,
        limit: int = 50,
        offset: int = 0,
    ) -> list[Document]:
        """List user's documents."""
        result = await self.db.execute(
            select(Document)
            .where(Document.user_id == user_id)
            .order_by(Document.created_at.desc())
            .offset(offset)
            .limit(limit)
        )
        return list(result.scalars().all())

    async def delete_document(self, document: Document) -> bool:
        """Delete a document and its file."""
        # Delete file
        file_path = Path(document.file_path)
        if file_path.exists():
            file_path.unlink()

        # Delete database record (cascades to chunks)
        await self.db.delete(document)
        await self.db.flush()

        return True
