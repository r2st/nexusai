"""Main RAG pipeline orchestrator."""

import time
from pathlib import Path
from typing import Any
from uuid import UUID

import structlog

from app.services.rag.config import RAGConfig, get_rag_config
from app.services.rag.models import (
    Document,
    DocumentChunk,
    DocumentMetadata,
    DocumentType,
    RAGQuery,
    RAGResult,
    RetrievalResult,
)
from app.services.rag.chunking import get_chunker
from app.services.rag.embedding import get_embedder
from app.services.rag.retrieval import Retriever, VectorStore

logger = structlog.get_logger()


class RAGPipeline:
    """
    Complete RAG pipeline for document processing and retrieval.

    Stages:
    1. Ingestion: Load and parse documents
    2. Chunking: Split into semantic chunks
    3. Embedding: Generate vector embeddings
    4. Storage: Store in vector database
    5. Retrieval: Search for relevant chunks
    6. Reranking: Re-score for relevance
    7. Context: Assemble context for LLM
    """

    def __init__(self, config: RAGConfig | None = None):
        self.config = config or get_rag_config()
        self.embedder = get_embedder(self.config.embedding)
        self.vector_store = VectorStore(self.config.retrieval)
        self.retriever = Retriever(self.config.retrieval, self.vector_store)

    async def ingest_document(
        self,
        content: str,
        metadata: DocumentMetadata | None = None,
        doc_type: DocumentType = DocumentType.TEXT,
    ) -> Document:
        """
        Ingest a document into the RAG system.

        Args:
            content: Document content
            metadata: Document metadata
            doc_type: Type of document

        Returns:
            The indexed document
        """
        from uuid import uuid4

        metadata = metadata or DocumentMetadata()
        metadata.source_type = doc_type

        # Create document
        document = Document(
            id=uuid4(),
            content=content,
            metadata=metadata,
        )

        try:
            # Chunk document
            chunker = get_chunker(doc_type, self.config.chunking)
            chunks = chunker.chunk(
                content,
                metadata={"document_id": document.id, **metadata.dict()},
            )

            logger.info(
                "Chunked document",
                document_id=str(document.id),
                chunk_count=len(chunks),
            )

            # Generate embeddings
            chunk_texts = [c.content for c in chunks]
            embeddings = await self.embedder.embed(chunk_texts)

            # Attach embeddings to chunks
            for chunk, embedding in zip(chunks, embeddings):
                chunk.embedding = embedding

            # Store in vector database
            await self.vector_store.add_chunks(chunks)

            # Update document stats
            document.chunk_count = len(chunks)
            document.token_count = sum(c.token_count for c in chunks)
            document.indexed = True

            logger.info(
                "Ingested document",
                document_id=str(document.id),
                chunks=len(chunks),
                tokens=document.token_count,
            )

            return document

        except Exception as e:
            document.indexed = False
            document.index_error = str(e)
            logger.error(
                "Failed to ingest document",
                document_id=str(document.id),
                error=str(e),
            )
            raise

    async def ingest_file(
        self,
        file_path: str | Path,
        metadata: DocumentMetadata | None = None,
    ) -> Document:
        """
        Ingest a file into the RAG system.

        Args:
            file_path: Path to the file
            metadata: Additional metadata

        Returns:
            The indexed document
        """
        path = Path(file_path)

        if not path.exists():
            raise FileNotFoundError(f"File not found: {file_path}")

        # Determine document type
        ext = path.suffix.lower()
        doc_type_map = {
            ".md": DocumentType.MARKDOWN,
            ".txt": DocumentType.TEXT,
            ".pdf": DocumentType.PDF,
            ".html": DocumentType.HTML,
            ".htm": DocumentType.HTML,
            ".py": DocumentType.CODE,
            ".js": DocumentType.CODE,
            ".ts": DocumentType.CODE,
            ".jsx": DocumentType.CODE,
            ".tsx": DocumentType.CODE,
            ".java": DocumentType.CODE,
            ".go": DocumentType.CODE,
            ".rs": DocumentType.CODE,
            ".cpp": DocumentType.CODE,
            ".c": DocumentType.CODE,
            ".docx": DocumentType.DOCX,
            ".csv": DocumentType.CSV,
            ".json": DocumentType.JSON,
        }
        doc_type = doc_type_map.get(ext, DocumentType.TEXT)

        # Load content based on type
        if doc_type == DocumentType.PDF:
            content = await self._load_pdf(path)
        elif doc_type == DocumentType.DOCX:
            content = await self._load_docx(path)
        else:
            content = path.read_text(encoding="utf-8")

        # Create metadata
        if metadata is None:
            metadata = DocumentMetadata()

        metadata.title = metadata.title or path.stem
        metadata.source = str(path)
        metadata.source_type = doc_type

        return await self.ingest_document(content, metadata, doc_type)

    async def _load_pdf(self, path: Path) -> str:
        """Load text from PDF."""
        try:
            import pypdf

            reader = pypdf.PdfReader(str(path))
            text = ""
            for page in reader.pages:
                text += page.extract_text() + "\n"
            return text
        except ImportError:
            raise ImportError("pypdf required for PDF processing: pip install pypdf")

    async def _load_docx(self, path: Path) -> str:
        """Load text from DOCX."""
        try:
            from docx import Document as DocxDocument

            doc = DocxDocument(str(path))
            return "\n\n".join(para.text for para in doc.paragraphs)
        except ImportError:
            raise ImportError("python-docx required for DOCX: pip install python-docx")

    async def query(self, query: RAGQuery) -> RAGResult:
        """
        Query the RAG system.

        Args:
            query: The RAG query

        Returns:
            Query results with context
        """
        start_time = time.time()

        # Generate query embedding
        query_embedding = await self.embedder.embed_query(query.query)

        # Retrieve relevant chunks
        retrieval_start = time.time()
        results = await self.retriever.retrieve(query, query_embedding)
        retrieval_time = (time.time() - retrieval_start) * 1000

        # Rerank if enabled
        rerank_time = 0.0
        if query.rerank and results and self.config.rerank.enabled:
            rerank_start = time.time()
            results = await self._rerank(query.query, results, query.rerank_top_k)
            rerank_time = (time.time() - rerank_start) * 1000

        # Assemble context
        context, sources = self._assemble_context(results)

        total_time = (time.time() - start_time) * 1000

        return RAGResult(
            query=query.query,
            results=results,
            context=context,
            context_token_count=len(context) // 4 if context else 0,
            retrieval_time_ms=retrieval_time,
            rerank_time_ms=rerank_time,
            total_time_ms=total_time,
            sources=sources,
        )

    async def _rerank(
        self,
        query: str,
        results: list[RetrievalResult],
        top_k: int,
    ) -> list[RetrievalResult]:
        """Rerank results using a cross-encoder or Cohere."""
        if self.config.rerank.provider == "cross-encoder":
            return await self._rerank_cross_encoder(query, results, top_k)
        elif self.config.rerank.provider == "cohere":
            return await self._rerank_cohere(query, results, top_k)
        return results[:top_k]

    async def _rerank_cross_encoder(
        self,
        query: str,
        results: list[RetrievalResult],
        top_k: int,
    ) -> list[RetrievalResult]:
        """Rerank using a cross-encoder model."""
        try:
            from sentence_transformers import CrossEncoder

            model = CrossEncoder(self.config.rerank.model)
            pairs = [(query, r.chunk.content) for r in results]
            scores = model.predict(pairs)

            for result, score in zip(results, scores):
                result.rerank_score = float(score)
                result.combined_score = (result.score + float(score)) / 2

            results.sort(key=lambda x: x.rerank_score or 0, reverse=True)
            return results[:top_k]

        except ImportError:
            logger.warning("sentence-transformers not installed, skipping rerank")
            return results[:top_k]

    async def _rerank_cohere(
        self,
        query: str,
        results: list[RetrievalResult],
        top_k: int,
    ) -> list[RetrievalResult]:
        """Rerank using Cohere's rerank API."""
        try:
            import httpx
            from app.core.config import settings

            if not getattr(settings, "COHERE_API_KEY", None):
                logger.warning("COHERE_API_KEY not set, skipping rerank")
                return results[:top_k]

            async with httpx.AsyncClient() as client:
                response = await client.post(
                    "https://api.cohere.ai/v1/rerank",
                    headers={
                        "Authorization": f"Bearer {settings.COHERE_API_KEY}",
                        "Content-Type": "application/json",
                    },
                    json={
                        "model": self.config.rerank.cohere_model,
                        "query": query,
                        "documents": [r.chunk.content for r in results],
                        "top_n": top_k,
                    },
                    timeout=30,
                )
                response.raise_for_status()
                data = response.json()

                # Update scores and reorder
                reranked = []
                for item in data["results"]:
                    idx = item["index"]
                    results[idx].rerank_score = item["relevance_score"]
                    results[idx].combined_score = (
                        results[idx].score + item["relevance_score"]
                    ) / 2
                    reranked.append(results[idx])

                return reranked

        except Exception as e:
            logger.warning("Cohere rerank failed", error=str(e))
            return results[:top_k]

    def _assemble_context(
        self,
        results: list[RetrievalResult],
    ) -> tuple[str, list[dict[str, Any]]]:
        """Assemble context from retrieval results."""
        if not results:
            return "", []

        context_parts = []
        sources = []
        seen_content = set()

        for i, result in enumerate(results):
            content = result.chunk.content.strip()

            # Deduplicate
            if self.config.context.deduplicate:
                content_hash = hash(content[:100])
                if content_hash in seen_content:
                    continue
                seen_content.add(content_hash)

            # Format context
            if self.config.context.include_metadata:
                source = result.document_source or "Unknown"
                context_parts.append(f"[Source {i+1}: {source}]\n{content}")
            else:
                context_parts.append(content)

            # Track sources
            sources.append({
                "index": i + 1,
                "chunk_id": str(result.chunk.id),
                "document_id": str(result.chunk.document_id),
                "title": result.document_title,
                "source": result.document_source,
                "score": result.score,
            })

        context = "\n\n---\n\n".join(context_parts)

        # Trim to max tokens
        max_chars = self.config.context.max_tokens * 4
        if len(context) > max_chars:
            context = context[:max_chars] + "..."

        return context, sources

    async def delete_document(self, document_id: UUID) -> bool:
        """Delete a document and its chunks."""
        deleted = await self.vector_store.delete_by_document(document_id)
        logger.info("Deleted document", document_id=str(document_id), chunks=deleted)
        return deleted > 0


# Global pipeline instance
_rag_pipeline: RAGPipeline | None = None


def get_rag_pipeline() -> RAGPipeline:
    """Get the global RAG pipeline."""
    global _rag_pipeline
    if _rag_pipeline is None:
        _rag_pipeline = RAGPipeline()
    return _rag_pipeline
